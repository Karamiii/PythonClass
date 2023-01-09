import openpyxl
from datetime import datetime, timedelta

# Create a new workbook and a worksheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sheet1"

# Set the starting date and the number of years to go forward
start_date = datetime(2020, 1, 1)
num_years = 5

# Set the table style
table_style = openpyxl.worksheet.table.TableStyleInfo(
    name='TableStyleMedium9',
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False
)

# Use a for loop to fill the cells with the dates
row = 1
for i in range(num_years * 365):
    # Calculate the current date
    current_date = start_date + timedelta(days=i)

    # Check if the current date is a Saturday
    if current_date.weekday() == 5:
        # Convert the date to a string and fill the cell
        ws.cell(row=row, column=1).value = current_date.strftime('%Y-%m-%d')
        row += 1

# Set the column width and create the table
ws.column_dimensions['A'].width = 15
table = openpyxl.worksheet.table.Table(ref='A1:A{}'.format(row - 1), displayName='Table1', tableStyleInfo=table_style)
ws.add_table(table)

# Save the workbook
wb.save('workbook.xlsx')
